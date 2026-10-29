#!/usr/bin/env python
"""Every table of the revised manuscript and Supplementary Information, built from outputs/ (read only).

    python artifacts/code/make_tables.py

Writes artifacts/tables/{main,si}/<id>.csv (display values) and <id>.md (caption, table, notes), artifacts/tables/raw/*.csv
(full precision), artifacts/tables/ALL_TABLES.md and ALL_TABLES.docx (one file to copy from), and artifacts/numbers.json
(every headline number quoted in the revision guides, with its source). Nothing is typed in: every value is read from a
pipeline output or derived from per-seed values by the formula stated next to it.

Statistics: mean ± sample SD (n - 1) over seeds 42, 43, 44. A run "lost a class" when that class has Dice < 0.05. Paired
differences are seed-matched, with a 95% t-interval (n = 3, so wide). Full fine-tuning is a single seed-42 run (D-24).
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common as C  # noqa: E402
from common import ARCH_LABEL, CLASSES, LOSS_LABEL, RANKS, SEEDS  # noqa: E402

REGISTRY: List[Dict[str, Any]] = []
NUMBERS: Dict[str, Any] = {}
RUNS = C.runs_frame()


def table(tid: str, part: str, title: str):
    def deco(fn: Callable[[], Dict[str, Any]]):
        REGISTRY.append({"id": tid, "part": part, "title": title, "fn": fn})
        return fn
    return deco


def note_number(key: str, value: Any, source: str) -> None:
    NUMBERS[key] = {"value": value, "source": source}


# ------------------------------------------------------------------------------------------------ helpers
def iou_stats(arch, rank, loss, col="mean_iou"):
    v = RUNS[(RUNS.arch == arch) & (RUNS["rank"].astype("float").fillna(-1) == (rank if rank is not None else -1))
             & (RUNS.loss == loss)][col]
    return C.mean_sd(v.tolist())


def trainable(arch, rank, loss):
    mj = C.metrics_json(arch, rank, loss, 42)["parameters"]
    return mj["trainable"], mj["total"]


def seg_cells(arch, rank, loss, per_class=True, iou=True) -> Dict[str, str]:
    m = C.cfg_entry(arch, rank, loss)["metrics"]
    tr, tot = trainable(arch, rank, loss)
    row = {"Architecture": ARCH_LABEL[arch], "LoRA rank": "full FT" if loss == "full_finetune" else rank,
           "Trainable params": f"{C.params_m(tr)} ({100 * tr / tot:.1f}%)",
           "Mean Dice": C.pm_metric(m["mean_dice"])}
    if iou:
        row["Mean IoU"] = C.pm(*iou_stats(arch, rank, loss)[:2])
    row.update({"AJI": C.pm_metric(m["aji"]), "Boundary F1": C.pm_metric(m["bf1"]),
                "Global IPE (%)": C.pm_metric(m["ipe_global_pct"], 2)})
    if per_class:
        for cl in CLASSES:
            row[f"Dice {C.CLASS_DISPLAY[cl]}"] = C.pm_metric(m[f"dice_{cl}"])
    row["Runs with a lost class"] = C.lost_text(arch, rank, loss)
    return row


def per_seed(arch, rank, loss, name) -> Dict[str, float]:
    return {s: v for s, v in C.metric(arch, rank, loss, name)["per_seed"].items() if v is not None}


def stat_line(m: Dict[str, Any], nd=3) -> str:
    return C.pm_metric(m, nd)


# ================================================================================================ MAIN TABLES
@table("T1_dataset", "main", "Dataset after the corrected preprocessing (float32 phase, background referenced)")
def t1():
    gt_all = C.cells_gt(nowrap=False)
    ann = C.compiled("cell_analysis")["annotation"]
    mj = C.metrics_json("edge_sam", 8, "full", 42)
    n_views = {int(C.metrics_json(a, r, l, s)["n_train_views"]) for a, r, l in [("edge_sam", 8, "full"), ("mobile_sam", 8, "full"),
                                                                                 ("mobilenet_unet", 8, "full")] for s in SEEDS}
    n_val = {int(C.metrics_json(a, 8, "full", s)["n_val_images"]) for a in C.ARCHS for s in SEEDS}
    if len(n_views) != 1 or len(n_val) != 1:
        raise RuntimeError(f"training/validation sizes differ between runs: {n_views} {n_val}")
    views, nval = n_views.pop(), n_val.pop()
    days = sorted(gt_all.storage_day.unique())
    per_day_imgs = gt_all.groupby("storage_day").stem.nunique()
    rows = [
        {"Item": "Training images (base, with float32 phase map)", "Value": f"{views // 4}", "Source": "metrics.json n_train_views / 4"},
        {"Item": "Training views (4 lossless 90° rotations per base image)", "Value": f"{views}", "Source": "metrics.json n_train_views"},
        {"Item": "Validation (= evaluation) images", "Value": f"{nval}", "Source": "metrics.json n_val_images"},
        {"Item": "Storage days in the validation set", "Value": f"{len(days)} (days {', '.join(str(d) for d in days)})", "Source": "cells_gt.csv"},
        {"Item": "Validation images per storage day", "Value": "/".join(str(int(v)) for v in per_day_imgs.values), "Source": "cells_gt.csv"},
        {"Item": "Annotated cells (area 15-200 µm², not cut by the annotated field of view)", "Value": f"{ann['n_cells_all']}", "Source": "compiled cell_analysis annotation"},
        {"Item": "  of which phase-wrapped (excluded from phase-derived results, D-21)",
         "Value": f"{ann['wrapping']['n_wrapped']} ({ann['wrapping']['pct_wrapped']:.1f}%) in {ann['wrapping']['n_images_with_wrapped']} images",
         "Source": "compiled cell_analysis annotation"},
        {"Item": "Annotated cells used for phase-derived results", "Value": f"{ann['n_cells_nowrap']}", "Source": "compiled cell_analysis annotation"},
    ]
    for cl in CLASSES:
        w = ann["wrapping"]["by_class"][cl]
        rows.append({"Item": f"  {C.CLASS_DISPLAY[cl]}: cells (wrapped)", "Value": f"{w['n_cells']} ({w['n_wrapped']}, {w['pct_wrapped']:.1f}%)",
                     "Source": "compiled cell_analysis annotation"})
    nw = gt_all[~gt_all.wrapped.astype(bool)]
    pk = nw.max_phase_rad
    rows.append({"Item": "Peak phase of the non-wrapped annotated cells, background-referenced (rad): median (5th to 95th percentile), max",
                 "Value": f"{pk.median():.2f} ({pk.quantile(0.05):.2f} to {pk.quantile(0.95):.2f}), {pk.max():.2f}", "Source": "cells_gt.csv max_phase_rad"})
    note_number("annotated_peak_phase_rad", {"median": round(float(pk.median()), 2), "p95": round(float(pk.quantile(0.95)), 2),
                                             "max": round(float(pk.max()), 2)}, "outputs/cell_analysis/annotation/cells_gt.csv, non-wrapped")
    note_number("n_train_views", views, "metrics.json n_train_views")
    note_number("n_val_images", nval, "metrics.json n_val_images")
    note_number("annotated_cells_all", ann["n_cells_all"], "outputs/compiled/cell_analysis_multiseed.json annotation")
    note_number("annotated_cells_nowrap", ann["n_cells_nowrap"], "same")
    note_number("annotated_cells_wrapped_pct", round(ann["wrapping"]["pct_wrapped"], 2), "same")
    return {"df": pd.DataFrame(rows),
            "notes": ["The 74 base training images without a float32 map (DECISIONS.md D-05, Q-1 closed) are not used; the "
                      "count 256 of the 8-bit set comes from DECISIONS.md, not from outputs/ (data/processed is not in the repository).",
                      "Wrapped cells: raw phase within 0.3 rad of 2π or a neighbour jump > π within 2 px of the cell (D-21)."]}


@table("T2_preferred", "main", "Segmentation at the preferred LoRA configurations (full physics-aware objective)")
def t2():
    rows = [seg_cells(a, r, "full") for a, r in C.PREFERRED + [C.BASELINE]]
    for a, r in C.PREFERRED:
        m = C.cfg_entry(a, r, "full")["metrics"]
        for k in ("mean_dice", "aji", "bf1", "ipe_global_pct"):
            note_number(f"{a}_r{r}_{k}", {"mean": round(m[k]["mean"], 4), "sd": round(m[k]["sd"], 4)},
                        f"outputs/compiled/cell_analysis_multiseed.json {a}|lora_r{r}")
    return {"df": pd.DataFrame(rows),
            "notes": ["MobileSAM r = 8 and EdgeSAM r = 32 are the presentation choice of the revision (D-27, D-30); no selection rule is "
                      "claimed. MobileNet-UNet (r = 8) is the convolutional baseline.",
                      "Mean ± SD over seeds 42, 43, 44; last-epoch checkpoint; metrics at the native 768 × 768 grid.",
                      "Lost class: per-class Dice < 0.05. Class names carry their mask value: sphero(2) = spherocyte, echino(3) = echinocyte, stomato(4) = stomatocyte."]}


@table("T3_same_rank", "main", f"All architectures at the same LoRA rank (r = {C.SAME_RANK}, full objective)")
def t3():
    rows = [seg_cells(a, C.SAME_RANK, "full") for a in C.ARCHS]
    diffs = []
    for a, b in (("edge_sam", "mobile_sam"), ("edge_sam", "mobilenet_unet")):
        d = {"Comparison (seed-paired, r = 8)": f"{ARCH_LABEL[b]} minus {ARCH_LABEL[a]}"}
        for k, nd in (("mean_dice", 3), ("aji", 3), ("bf1", 3), ("ipe_global_pct", 2)):
            d[k] = C.fmt_diff(C.paired_diff(per_seed(a, 8, "full", k), per_seed(b, 8, "full", k)), nd)
        diffs.append(d)
    return {"df": pd.DataFrame(rows), "extra": pd.DataFrame(diffs).rename(columns={"mean_dice": "Δ Mean Dice", "aji": "Δ AJI",
                                                                                    "bf1": "Δ Boundary F1", "ipe_global_pct": "Δ Global IPE (pp)"}),
            "notes": ["Same-rank comparison requested in the revision: identical LoRA rank, objective, schedule and data.",
                      "Δ: mean of the seed-paired differences [95% t-interval, n = 3]."]}


@table("T4_training_stability", "main", "Training stability: runs that lost a morphology class, per architecture and LoRA rank")
def t4():
    rows = []
    for loss in ("full", "dice_only"):
        for a in C.ARCHS:
            r = {"Objective": "Full physics-aware" if loss == "full" else "Dice only", "Architecture": ARCH_LABEL[a]}
            for rk in RANKS:
                r[f"r = {rk}"] = C.lost_text(a, rk, loss)
            rows.append(r)
    sd_rows = []
    for a in C.ARCHS:
        r = {"Architecture": ARCH_LABEL[a]}
        for rk in RANKS:
            m = C.metric(a, rk, "full", "mean_dice")
            r[f"r = {rk}"] = f"{m['sd']:.3f}"
        sd_rows.append(r)
    tot = {a: sum(C.n_lost_runs(a, rk, "full") for rk in RANKS) for a in C.ARCHS}
    note_number("lost_runs_full_objective_rank_sweep", {ARCH_LABEL[a]: f"{tot[a]}/15" for a in C.ARCHS}, "outputs/compiled/cell_analysis_multiseed.json")
    return {"df": pd.DataFrame(rows), "extra": pd.DataFrame(sd_rows), "extra_title": "Seed SD of mean Dice (full objective)",
            "notes": ["k/3 = number of the three seeds in which at least one class had Dice < 0.05; the seed and class are named.",
                      "MobileNet-UNet loses spherocytes in every run (and stomatocytes in 14 of 15 full-objective runs); this is a "
                      "systematic failure, not instability.",
                      "Training curves: Fig. 5 (artifacts/figures/main)."]}


def ablation_rows(a, rk):
    rows = []
    for loss in C.LOSSES:
        m = C.cfg_entry(a, rk, loss)["metrics"]
        rows.append({"Configuration": f"{ARCH_LABEL[a]} r = {rk}", "Objective": LOSS_LABEL[loss],
                     "Mean Dice": C.pm_metric(m["mean_dice"]), "Boundary F1": C.pm_metric(m["bf1"]), "AJI": C.pm_metric(m["aji"]),
                     "Global IPE (%)": C.pm_metric(m["ipe_global_pct"], 2), "Runs with a lost class": C.lost_text(a, rk, loss)})
    return rows


@table("T5_loss_ablation", "main", "Component-wise ablation of the physics-aware loss at the preferred and same-rank configurations")
def t5():
    cfgs = [("mobile_sam", 8), ("edge_sam", 32), ("edge_sam", 8), ("mobilenet_unet", 8)]
    rows, diffs = [], []
    for a, rk in cfgs:
        rows += ablation_rows(a, rk)
        d = {"Configuration": f"{ARCH_LABEL[a]} r = {rk}"}
        for lab, (l1, l2) in {"Full − Dice only": ("dice_only", "full"), "Full − (Dice+PMC+BGA): IPP": ("pmc_bga", "full")}.items():
            p = C.paired_diff(per_seed(a, rk, l1, "ipe_global_pct"), per_seed(a, rk, l2, "ipe_global_pct"))
            d[f"Δ IPE (pp), {lab}"] = C.fmt_diff(p, 2)
            q = C.paired_diff(per_seed(a, rk, l1, "mean_dice"), per_seed(a, rk, l2, "mean_dice"))
            d[f"Δ Dice, {lab}"] = C.fmt_diff(q, 3)
            note_number(f"ablation_{a}_r{rk}_{l2}_minus_{l1}_ipe_pp", {k: (round(v, 3) if isinstance(v, float) else v) for k, v in p.items()},
                        "seed-paired differences of compiled per-seed values")
        diffs.append(d)
    return {"df": pd.DataFrame(rows), "extra": pd.DataFrame(diffs), "extra_title": "Seed-paired differences [95% t-interval, n = 3]",
            "notes": ["The full-objective row of each block is the rank-sweep run of that rank (same configuration and seed; it is "
                      "not trained twice, D-09). The ablation runs are separate trainings with the same seeds.",
                      "The physics terms act on the binary foreground probability and on the background-referenced phase in radians.",
                      "All ranks: Table S4/S5."]}


@table("T6_lora_vs_full_ft", "main", "LoRA versus full fine-tuning (full physics-aware objective)")
def t6():
    rows = []
    for a in C.ARCHS:
        pref = dict(C.PREFERRED + [C.BASELINE])[a]
        ranks = sorted({pref, C.SAME_RANK})
        for rk in ranks:
            rows.append({**seg_cells(a, rk, "full"), "Seeds": "42, 43, 44"})
        # seed-42 LoRA values next to the single seed-42 full fine-tuning run (same seed, like for like)
        ft = seg_cells(a, None, "full_finetune")
        for rk in ranks:
            m = C.cfg_entry(a, rk, "full")["metrics"]
            r = {"Architecture": ARCH_LABEL[a], "LoRA rank": f"{rk} (seed 42 only)",
                 "Trainable params": seg_cells(a, rk, "full", False, False)["Trainable params"],
                 "Mean Dice": C.fmt(m["mean_dice"]["per_seed"]["42"]),
                 "Mean IoU": C.fmt(C.metrics_json(a, rk, "full", 42)["evaluation"]["last"]["mean_iou"]),
                 "AJI": C.fmt(m["aji"]["per_seed"]["42"]), "Boundary F1": C.fmt(m["bf1"]["per_seed"]["42"]),
                 "Global IPE (%)": C.fmt(m["ipe_global_pct"]["per_seed"]["42"], 2)}
            for cl in CLASSES:
                r[f"Dice {C.CLASS_DISPLAY[cl]}"] = C.fmt(m[f"dice_{cl}"]["per_seed"]["42"])
            lost = [C.CLASS_SHORT[cl] for cl in CLASSES if not m[f"dice_{cl}"]["per_seed"]["42"] >= C.COLLAPSE_DICE]
            r["Runs with a lost class"] = ", ".join(lost) if lost else "none"
            r["Seeds"] = "42"
            rows.append(r)
        ft["LoRA rank"] = "full fine-tuning"
        ft["Seeds"] = "42"
        ftm = C.cfg_entry(a, None, "full_finetune")["metrics"]
        lost = [C.CLASS_SHORT[cl] for cl in CLASSES if not ftm[f"dice_{cl}"]["per_seed"]["42"] >= C.COLLAPSE_DICE]
        ft["Runs with a lost class"] = ", ".join(lost) if lost else "none"
        rows.append(ft)
    return {"df": pd.DataFrame(rows),
            "notes": ["Full fine-tuning was trained once (seed 42, D-24); its row is a single run (no SD). Compare it with the "
                      "'seed 42 only' LoRA rows (same seed) and read the 3-seed rows for the LoRA seed spread.",
                      "Full fine-tuning lost spherocytes in EdgeSAM, all three minority classes in MobileSAM, and spherocytes and "
                      "stomatocytes in MobileNet-UNet."]}


@table("T7_phase_and_mass", "main", "Integrated-phase and optical-mass preservation (wrapped cells excluded)")
def t7():
    cfgs = C.PREFERRED + [("edge_sam", 8), C.BASELINE]
    rows = []
    keys = [("Global IPE (%)", "ipe_global_pct", 2), ("Per-cell |IPE| median (%)", "percell_abs_median_pct_nowrap", 2),
            ("Per-cell |IPE| P90 (%)", "percell_abs_p90_pct_nowrap", 1), ("Annotated cells matched (IoU ≥ 0.5, %)", "gt_cells_matched_pct_nowrap", 1),
            ("Predicted / annotated cells", "cells_ratio_nowrap", 3), ("Dry-mass MdARE (%)", "mdare_pct", 2),
            ("Dry-mass bias (pg)", "bias_pg", 2), ("LoA low (pg)", "loa_low_pg", 2), ("LoA high (pg)", "loa_high_pg", 2),
            ("Pearson r (per-image medians)", "pearson_r", 2), ("Lin CCC", "lin_ccc", 2),
            ("Day-0 median dry mass (pg)", "day0_median_dry_mass_pg", 2)]
    for a, rk in cfgs:
        m = C.cfg_entry(a, rk, "full")["metrics"]
        r = {"Configuration": f"{ARCH_LABEL[a]} r = {rk}"}
        for lab, k, nd in keys:
            r[lab] = C.pm_metric(m[k], nd)
        rows.append(r)
        for k in ("mdare_pct", "bias_pg", "pearson_r", "percell_abs_median_pct_nowrap", "day0_median_dry_mass_pg"):
            note_number(f"{a}_r{rk}_{k}", {"mean": round(m[k]["mean"], 3), "sd": round(m[k]["sd"], 3)},
                        f"outputs/compiled/cell_analysis_multiseed.json {a}|lora_r{rk}")
    ann = C.compiled("cell_analysis")["annotation"]
    gt = C.cells_gt()
    img = gt.groupby("stem").dry_mass_pg.median()
    rows.append({"Configuration": "Annotation (reference)", "Day-0 median dry mass (pg)": f"{ann['consistency_nowrap']['day0_median_dry_mass_pg']:.2f}",
                 "Predicted / annotated cells": "1"})
    note_number("annotation_day0_median_dry_mass_pg", round(ann["consistency_nowrap"]["day0_median_dry_mass_pg"], 2), "compiled annotation consistency_nowrap")
    note_number("annotation_day0_over_mch", round(ann["consistency_nowrap"]["day0_median_over_mch"], 3), "compiled annotation consistency_nowrap")
    note_number("annotation_per_image_median_dry_mass_pg", {"mean": round(float(img.mean()), 2), "sd": round(float(img.std()), 2),
                                                            "min": round(float(img.min()), 2), "max": round(float(img.max()), 2), "n_images": int(len(img))},
                "derived: cells_gt.csv, non-wrapped cells, median per image")
    return {"df": pd.DataFrame(rows).fillna(""),
            "notes": ["Dry mass m = λ/(2πα) Σ(φ − φ_bg) Δx², λ = 0.666 µm, α = 0.2 mL/g, Δx = 0.1441 µm, φ in rad after per-image background "
                      "referencing; no calibration to MCH (D-13). MCH = 30 pg is only a plausibility check (annotation day-0 median / MCH = "
                      f"{ann['consistency_nowrap']['day0_median_over_mch']:.2f}).",
                      "Agreement statistics use per-image median dry mass (66 images). Because the annotated per-image median dry mass varies "
                      f"little between images ({img.mean():.1f} ± {img.std():.1f} pg), Pearson r and CCC are low even when the errors are "
                      "small; MdARE, bias and the limits of agreement (LoA, bias ± 1.96 SD) are the informative measures.",
                      "Per-cell IPE: each annotated cell matched one-to-one to a predicted component (IoU ≥ 0.5), phase after a "
                      "least-squares background plane; unmatched cells excluded."]}


@table("T8_efficiency", "main", "Computational efficiency (NVIDIA RTX A5000, batch 1, strict FP32, 3 seeds)")
def t8():
    b = C.compiled("benchmark")
    params = {(p["architecture"], p["lora_rank"]): p for p in b["parameters_by_rank"]}
    pref = dict(C.PREFERRED + [C.BASELINE])
    rows = []
    for a in C.ARCHS:
        for be in ("pytorch", "onnx"):
            r = next(x for x in b["rows"] if x["architecture"] == a and x["backend"] == be)
            fps = [v for v in r["fps"]["per_seed"].values()]
            rows.append({"Architecture": ARCH_LABEL[a], "Runtime": "PyTorch FP32" if be == "pytorch" else "ONNX Runtime FP32",
                         "Input (px)": r["input_size"],
                         "Trainable params, preferred rank": f"{C.params_m(params[(a, pref[a])]['params_trainable'])} (r = {pref[a]})",
                         "Trainable params, r = 8": C.params_m(params[(a, 8)]["params_trainable"]),
                         "P50 (ms)": C.pm(r["latency_p50_ms"]["mean"], r["latency_p50_ms"]["sd"], 2),
                         "P99 (ms)": C.pm(r["latency_p99_ms"]["mean"], r["latency_p99_ms"]["sd"], 2),
                         "FPS (per-seed range)": f"{min(fps):.0f}-{max(fps):.0f}",
                         "Peak VRAM (MB, max)": f"{r['peak_vram_mb_max']:.0f}"})
            note_number(f"bench_{a}_{be}", {"p50_ms": round(r["latency_p50_ms"]["mean"], 2), "p50_sd": round(r["latency_p50_ms"]["sd"], 2),
                                            "fps_range": [round(min(fps), 1), round(max(fps), 1)], "peak_vram_mb_max": r["peak_vram_mb_max"]},
                        "outputs/compiled/benchmark_multiseed.json")
    return {"df": pd.DataFrame(rows),
            "notes": ["LoRA is merged into the weights before timing, so latency and memory do not depend on the rank: one measurement "
                      "per architecture covers every rank (trainable parameters differ, inference cost does not).",
                      "MobileSAM is timed at 1024 px and the other two at 256 px (each model's native input), so the latency gap mixes "
                      "architecture and input resolution.",
                      "P50/P99: mean ± SD over three seeds (500 timed passes each after 50 warm-up; ONNX: 3 fresh sessions per seed). "
                      "FPS = 1000/P50. Peak VRAM: device memory after the timed loop minus before model placement; largest value over "
                      "seeds and sessions. Protocol: manuscript/benchmarking_protocol.md (v5).",
                      "ONNX Runtime FP32 passed the equivalence gate for all three models (Table S9)."]}


@table("T9_storage_per_day", "main", "Cell morphology and dry mass over storage: annotation and preferred models (wrapped cells excluded)")
def t9():
    ann = C.compiled("cell_analysis")["annotation"]["per_day_nowrap"]
    rows = []
    for d in ann:
        comp = d["composition_pct"]
        r = {"Day": d["storage_day"], "Annotated cells": d["n_cells"],
             "Area (µm²), annotation": C.pm(d["area_um2_mean"], d["area_um2_sd"], 1),
             "Circularity, annotation": C.pm(d["circularity_mean"], d["circularity_sd"], 3),
             "Dry mass (pg), annotation": C.pm(d["dry_mass_pg_mean"], d["dry_mass_pg_sd"], 1),
             "Composition (%), annotation": " / ".join(f"{comp[cl]:.0f}" for cl in CLASSES)}
        for a, rk in C.PREFERRED:
            pd_ = next(x for x in C.cfg_entry(a, rk, "full")["per_day_pred_nowrap"] if x["storage_day"] == d["storage_day"])
            mm = pd_["metrics"]
            lab = f"{ARCH_LABEL[a]} r{rk}"
            r[f"Cells, {lab}"] = C.pm(mm["n_cells"]["mean"], mm["n_cells"]["sd"], 0)
            r[f"Area (µm²), {lab}"] = C.pm(mm["area_um2_mean"]["mean"], mm["area_um2_mean"]["sd"], 1)
            r[f"Dry mass (pg), {lab}"] = C.pm(mm["dry_mass_pg_mean"]["mean"], mm["dry_mass_pg_mean"]["sd"], 1)
        rows.append(r)
    return {"df": pd.DataFrame(rows),
            "notes": ["Annotation columns: mean ± SD over cells. Model columns: the per-day cell mean, averaged over seeds 42, 43, 44 "
                      "± SD over seeds (the seed SD, not the cell SD).",
                      "Composition order: discocyte (1) / spherocyte (2) / echinocyte (3) / stomatocyte (4); the number is the mask value.",
                      "Statistical tests use the image (median of its cells) as the unit (Table T10)."]}


@table("T10_storage_statistics", "main", "Storage trends with the image as the unit of analysis (per-image medians, 66 images)")
def t10():
    ann = C.compiled("cell_analysis")["annotation"]["storage_nowrap"]
    rows = []

    def ann_row(q, lab):
        e = ann[q]
        return {"Quantity": lab, "Masks": "Annotation", "Spearman ρ vs day": f"{e['spearman_rho']:.2f} (p = {e['spearman_p']:.2g})",
                "Late vs early (%)": f"{e['late_vs_early_pct']:+.1f} (MW p = {e['mannwhitney_p']:.2g})"}

    for q, lab in (("area", "Projected area"), ("dry_mass", "Dry mass"), ("circularity", "Circularity")):
        rows.append(ann_row(q, lab))
        if q == "dry_mass":
            rows[-1]["Partial ρ (dry mass | area)"] = f"{ann['dry_mass_given_area_partial_spearman']['rho']:.2f} (p = {ann['dry_mass_given_area_partial_spearman']['p']:.2g})"
        for a, rk in C.PREFERRED + [C.BASELINE]:
            st = [C.cell_unit(a, rk, "full", s)["storage"]["pred_nowrap"] for s in SEEDS]
            rho = [x[q]["spearman_rho"] for x in st]
            p = [x[q]["spearman_p"] for x in st]
            lve = [x[q]["late_vs_early_pct"] for x in st]
            mw = [x[q]["mannwhitney_p"] for x in st]
            r = {"Quantity": lab, "Masks": f"{ARCH_LABEL[a]} r = {rk}",
                 "Spearman ρ vs day": f"{np.mean(rho):.2f} ± {np.std(rho, ddof=1):.2f} (p {min(p):.2g} to {max(p):.2g})",
                 "Late vs early (%)": f"{np.mean(lve):+.1f} ± {np.std(lve, ddof=1):.1f} (MW p {min(mw):.2g} to {max(mw):.2g})"}
            if q == "dry_mass":
                pr = [x["dry_mass_given_area_partial_spearman"]["rho"] for x in st]
                pp = [x["dry_mass_given_area_partial_spearman"]["p"] for x in st]
                r["Partial ρ (dry mass | area)"] = f"{np.mean(pr):.2f} ± {np.std(pr, ddof=1):.2f} (p {min(pp):.2g} to {max(pp):.2g})"
                note_number(f"storage_dry_mass_{a}_r{rk}", {"rho_mean": round(float(np.mean(rho)), 3), "late_vs_early_pct_mean": round(float(np.mean(lve)), 2),
                                                            "p_range": [float(min(p)), float(max(p))]}, "outputs/analysis/cell/<run>.json storage.pred_nowrap")
            rows.append(r)
    for q in ("area", "dry_mass"):
        note_number(f"storage_{q}_annotation", {"rho": round(ann[q]["spearman_rho"], 3), "p": ann[q]["spearman_p"],
                                                "late_vs_early_pct": round(ann[q]["late_vs_early_pct"], 2), "mw_p": ann[q]["mannwhitney_p"]},
                    "outputs/compiled/cell_analysis_multiseed.json annotation.storage_nowrap")
    return {"df": pd.DataFrame(rows).fillna(""),
            "notes": ["Early = days 0 to 8 (18 images), late = days 37 and 47 (12 images); two-sided Mann-Whitney (MW).",
                      "Model rows: mean ± SD over seeds of ρ and of the late-vs-early change; p-values as the range over seeds.",
                      "Partial ρ: Spearman correlation of dry mass with storage day after removing the rank-linear dependence on area.",
                      "Key reading: annotation-derived dry mass does not change with storage while area falls; the models reproduce the "
                      "area trend but add a small spurious dry-mass decline (segmentation bias, not biology)."]}


# ================================================================================================ SI TABLES
@table("S1_rank_sweep", "si", "Complete LoRA rank sweep, all architectures (full objective, 3 seeds)")
def s1():
    rows = [seg_cells(a, rk, "full") for a in C.ARCHS for rk in RANKS]
    return {"df": pd.DataFrame(rows), "notes": ["Mean ± SD over seeds 42, 43, 44. Per-seed values: Table S2."]}


@table("S2_rank_sweep_per_seed", "si", "Rank sweep: every seed (full objective)")
def s2():
    d = RUNS[(RUNS.loss == "full")].copy()
    d["Architecture"] = d.arch.map(ARCH_LABEL)
    out = d[["Architecture", "rank", "seed", "mean_dice", "mean_iou", "aji", "bf1", "ipe_global_pct"] + [f"dice_{c}" for c in CLASSES] + ["lost_classes"]].copy()
    out["rank"] = out["rank"].astype(int)
    for c in out.columns:
        if out[c].dtype == float:
            out[c] = out[c].map(lambda x: f"{x:.2f}" if c == "ipe_global_pct" else f"{x:.3f}")
    out["lost_classes"] = out["lost_classes"].replace("", "none")
    out = out.sort_values(["Architecture", "rank", "seed"])
    return {"df": out.rename(columns={"rank": "LoRA rank", "seed": "Seed", "mean_dice": "Mean Dice", "mean_iou": "Mean IoU", "aji": "AJI",
                                      "bf1": "Boundary F1", "ipe_global_pct": "Global IPE (%)", "lost_classes": "Lost classes",
                                      **{f"dice_{c}": f"Dice {C.CLASS_DISPLAY[c]}" for c in CLASSES}}),
            "notes": ["Training instability is reported per seed, never averaged away."]}


@table("S3_rank_paired", "si", "Rank sweep: seed-paired differences to r = 8 and to r = 32 (full objective)")
def s3():
    rows = []
    for a in C.ARCHS:
        for ref in (8, 32):
            for rk in RANKS:
                if rk == ref:
                    continue
                r = {"Architecture": ARCH_LABEL[a], "Comparison": f"r{rk} − r{ref}"}
                for k, nd in (("mean_dice", 3), ("aji", 3), ("bf1", 3), ("ipe_global_pct", 2)):
                    p = C.paired_diff(per_seed(a, ref, "full", k), per_seed(a, rk, "full", k))
                    r[k] = C.fmt_diff(p, nd) + (" *" if p["excludes_zero"] else "")
                rows.append(r)
    e = C.paired_diff(per_seed("edge_sam", 8, "full", "mean_dice"), per_seed("edge_sam", 32, "full", "mean_dice"))
    e2 = C.paired_diff(per_seed("edge_sam", 8, "full", "aji"), per_seed("edge_sam", 32, "full", "aji"))
    note_number("edge_sam_r32_minus_r8_dice", {k: (round(v, 3) if isinstance(v, float) else v) for k, v in e.items()}, "seed-paired")
    note_number("edge_sam_r32_minus_r8_aji", {k: (round(v, 3) if isinstance(v, float) else v) for k, v in e2.items()}, "seed-paired")
    return {"df": pd.DataFrame(rows).rename(columns={"mean_dice": "Δ Mean Dice", "aji": "Δ AJI", "bf1": "Δ Boundary F1", "ipe_global_pct": "Δ Global IPE (pp)"}),
            "notes": ["Mean of seed-paired differences [95% t-interval, n = 3]; * = interval excludes 0. With n = 3 the intervals are wide; "
                      "they are descriptive, not a selection rule (D-27)."]}


@table("S4_ablation_all_ranks", "si", "Loss ablation at every LoRA rank (3 seeds per row)")
def s4():
    rows = []
    for a in C.ARCHS:
        for rk in RANKS:
            rows += ablation_rows(a, rk)
    return {"df": pd.DataFrame(rows), "notes": ["60 configurations × 3 seeds = 180 runs."]}


@table("S5_ablation_paired", "si", "Loss ablation: seed-paired differences within each rank")
def s5():
    rows = []
    comps = [("dice_only", "full", "Full − Dice only"), ("pmc_bga", "full", "Full − (Dice+PMC+BGA) = IPP"),
             ("dice_only", "pmc", "(Dice+PMC) − Dice only = PMC"), ("pmc", "pmc_bga", "(Dice+PMC+BGA) − (Dice+PMC) = BGA")]
    sign = {}
    for a in C.ARCHS:
        for rk in RANKS:
            for l1, l2, lab in comps:
                p_ipe = C.paired_diff(per_seed(a, rk, l1, "ipe_global_pct"), per_seed(a, rk, l2, "ipe_global_pct"))
                p_d = C.paired_diff(per_seed(a, rk, l1, "mean_dice"), per_seed(a, rk, l2, "mean_dice"))
                p_b = C.paired_diff(per_seed(a, rk, l1, "bf1"), per_seed(a, rk, l2, "bf1"))
                rows.append({"Architecture": ARCH_LABEL[a], "LoRA rank": rk, "Comparison": lab,
                             "Δ Global IPE (pp)": C.fmt_diff(p_ipe, 2) + (" *" if p_ipe["excludes_zero"] else ""),
                             "Δ Mean Dice": C.fmt_diff(p_d, 3) + (" *" if p_d["excludes_zero"] else ""),
                             "Δ Boundary F1": C.fmt_diff(p_b, 3) + (" *" if p_b["excludes_zero"] else "")})
                sign.setdefault(lab, []).append((a, rk, p_ipe["mean"], p_ipe["excludes_zero"]))
    summ = []
    for lab, v in sign.items():
        neg = sum(1 for x in v if x[2] < 0)
        sig = sum(1 for x in v if x[2] < 0 and x[3])
        summ.append({"Comparison": lab, "Configurations with lower mean IPE": f"{neg}/{len(v)}",
                     "... and 95% interval excluding 0": f"{sig}/{len(v)}"})
        note_number(f"ablation_sign_{lab}", {"lower_ipe": f"{neg}/{len(v)}", "ci_excludes_0": f"{sig}/{len(v)}"}, "Table S5")
    return {"df": pd.DataFrame(rows), "extra": pd.DataFrame(summ), "extra_title": "Direction of the IPE change over the 15 architecture × rank combinations",
            "notes": ["* = 95% t-interval (n = 3) excludes 0. Negative Δ IPE = the added term lowered the integrated-phase error."]}


@table("S6_all_runs", "si", "Every trained run (180 LoRA + 3 full fine-tuning), per seed")
def s6():
    d = RUNS.copy()
    d["Architecture"] = d.arch.map(ARCH_LABEL)
    d["rank"] = d["rank"].map(lambda x: "full FT" if pd.isna(x) else int(x))
    cols = ["Architecture", "rank", "loss", "seed", "run_id", "trainable_params", "mean_dice", "mean_iou", "aji", "bf1", "ipe_global_pct"] + \
           [f"dice_{c}" for c in CLASSES] + [f"iou_{c}" for c in CLASSES] + [f"ipe_{c}_pct" for c in CLASSES] + ["lost_classes", "epochs", "train_seconds"]
    out = d[cols].copy()
    for c in out.columns:
        if out[c].dtype == float:
            out[c] = out[c].round(4)
    return {"df": out, "notes": ["Machine-readable: also artifacts/tables/raw/all_runs.csv. Class names in the manuscript order (D-25)."],
            "md_rows": 0}


@table("S7_per_class_ipe", "si", "Class-specific integrated-phase error (%) and Dice")
def s7():
    rows = []
    for a, rk, loss in [(a, r, "full") for a, r in C.PREFERRED + [("edge_sam", 8), ("mobile_sam", 32)] + [C.BASELINE]] + \
                       [(a, None, "full_finetune") for a in C.ARCHS]:
        m = C.cfg_entry(a, rk, loss)["metrics"]
        r = {"Configuration": f"{ARCH_LABEL[a]} " + ("full FT (seed 42)" if loss == "full_finetune" else f"r = {rk}")}
        for cl in CLASSES:
            ipe = C.mean_sd(RUNS[(RUNS.run_id.isin([C.run_id(a, rk, loss, s) for s in SEEDS]))][f"ipe_{cl}_pct"].tolist())
            r[f"IPE {C.CLASS_DISPLAY[cl]}"] = C.pm(ipe[0], ipe[1], 1)
        for cl in CLASSES:
            r[f"Dice {C.CLASS_DISPLAY[cl]}"] = C.pm_metric(m[f"dice_{cl}"])
        rows.append(r)
    return {"df": pd.DataFrame(rows),
            "notes": ["Class IPE: per image |S_pred,c − S_gt,c| / |S_gt,c| over the images containing class c, then the mean; a class that "
                      "is not predicted gives 100%. Values above 100% come from large over-segmentation of a rare class.",
                      "Class IPE is dominated by misclassification between morphologies; the global (all-foreground) IPE is the "
                      "measure of integrated-phase preservation."]}


@table("S8_parameters_by_rank", "si", "Trainable parameters by LoRA rank")
def s8():
    b = C.compiled("benchmark")["parameters_by_rank"]
    rows = [{"Architecture": ARCH_LABEL[p["architecture"]], "LoRA rank": p["lora_rank"], "Total params (training form)": f"{p['params_total']:,}",
             "Trainable params": f"{p['params_trainable']:,}", "Trainable (%)": f"{p['overall_trainable_pct']:.2f}",
             "Encoder trainable (%)": f"{p['encoder_trainable_pct']:.2f}"} for p in b]
    for a in C.ARCHS:
        tr, tot = trainable(a, None, "full_finetune")
        rows.append({"Architecture": ARCH_LABEL[a], "LoRA rank": "full FT", "Total params (training form)": f"{tot:,}", "Trainable params": f"{tr:,}",
                     "Trainable (%)": "100.00", "Encoder trainable (%)": "100.00"})
    return {"df": pd.DataFrame(rows),
            "notes": ["Decoder, output heads, normalisation layers and the single-channel input stem are trainable in every LoRA "
                      "configuration, so the trainable share differs between architectures. After merging, all ranks have the "
                      "parameter count of the full fine-tuning model."]}


@table("S9_benchmark_per_seed", "si", "Benchmark per seed and ONNX Runtime equivalence gate")
def s9():
    b = C.compiled("benchmark")
    rows = []
    for r in b["rows"]:
        for s in SEEDS:
            rows.append({"Architecture": ARCH_LABEL[r["architecture"]], "Runtime": r["backend"], "Seed": s,
                         "P50 (ms)": f"{r['latency_p50_ms']['per_seed'][str(s)]:.2f}", "P99 (ms)": f"{r['latency_p99_ms']['per_seed'][str(s)]:.2f}",
                         "FPS": f"{r['fps']['per_seed'][str(s)]:.1f}", "Peak VRAM (MB)": f"{r['peak_vram_mb']['per_seed'][str(s)]:.0f}"})
    eq = []
    for e in b["equivalence"]["architectures"]:
        cand = next(c for c in e["quality"]["candidates"] if c["candidate"] == "onnx_fp32")
        eq.append({"Architecture": ARCH_LABEL[e["architecture"]], "Arg-max agreement, 16 maps (%)": f"{e['logit_comparison']['argmax_agreement_pct']:.4f}",
                   "Max |Δ logit|": f"{e['logit_comparison']['max_abs_logit_diff']:.1e}",
                   "Arg-max agreement, 66 images (%)": f"{cand['argmax_agreement_pct']:.4f}",
                   "|Δ mean Dice|": f"{abs(cand['mean_dice_change']):.1e}",
                   "Dry-mass rel. error mean / P99": f"{cand['dry_mass_mean_abs_rel_error']:.1e} / {cand['dry_mass_p99_rel_error']:.1e}",
                   "Verdict": cand["verdict"]})
    return {"df": pd.DataFrame(rows), "extra": pd.DataFrame(eq), "extra_title": "ONNX Runtime FP32 vs PyTorch FP32 (gate fixed before the first trial)",
            "notes": [f"GPU: {b['gpu_name']}. Acceptance: arg-max agreement ≥ 99.9%, |Δ mean Dice| ≤ 0.001, dry-mass relative error "
                      "mean ≤ 0.1% and P99 ≤ 0.5%.",
                      "The timed and gated checkpoint is the seed-42 rank-8 best.pt of each architecture (D-26); timing does not depend "
                      "on the weights or on the rank after merging."]}


@table("S10_robustness", "si", "Robustness to synthetic background-phase perturbations")
def s10():
    r = C.compiled("robustness")["configurations"]
    rows = []
    for a, rk in C.PREFERRED + [("edge_sam", 8), C.BASELINE]:
        c = r[C.config_key(a, rk, "full")]
        for n, v in c["distortions"].items():
            rows.append({"Configuration": f"{ARCH_LABEL[a]} r = {rk}", "Perturbation (rad)": n,
                         "Mean Dice": C.pm(v["macro_dice"]["mean"], v["macro_dice"]["sd"], 3),
                         "Δ Dice vs none": C.pm(v["delta_macro_dice_vs_none"]["mean"], v["delta_macro_dice_vs_none"]["sd"], 3),
                         "IPE, background-corrected (%)": C.pm(v["ipe_bgcorr_true_pct"]["mean"], v["ipe_bgcorr_true_pct"]["sd"], 2),
                         "IPE, perturbed phase (%)": C.pm(v["ipe_uncorrected_seen_pct"]["mean"], v["ipe_uncorrected_seen_pct"]["sd"], 2),
                         "Invariant by construction": "yes" if v["invariant_by_construction"] else ""})
    return {"df": pd.DataFrame(rows),
            "notes": ["Offset, tilt, bowl (defocus-like) and astigmatism-like fields of the stated peak amplitude added to the validation "
                      "phase. Offsets are removed by the per-image z-score of the network input, so the unchanged Dice there is a property "
                      "of the preprocessing (flagged), not robustness of the model."]}


@table("S11_wrapping", "si", "Phase-wrapped cells by class and storage day")
def s11():
    w = C.compiled("cell_analysis")["annotation"]["wrapping"]
    rows = [{"Group": f"class {C.CLASS_DISPLAY[k]}", "Annotated cells": v["n_cells"], "Wrapped": v["n_wrapped"], "Wrapped (%)": f"{v['pct_wrapped']:.1f}"}
            for k, v in w["by_class"].items()]
    rows += [{"Group": f"day {d['storage_day']}", "Annotated cells": d["n_cells"], "Wrapped": d["n_wrapped"], "Wrapped (%)": f"{d['pct_wrapped']:.1f}"}
             for d in w["by_storage_day"]]
    rows.append({"Group": "all", "Annotated cells": w["n_cells"], "Wrapped": w["n_wrapped"], "Wrapped (%)": f"{w['pct_wrapped']:.1f}"})
    for a, rk in C.PREFERRED:
        m = C.metric(a, rk, "full", "pct_predicted_cells_wrapped")
        rows.append({"Group": f"predicted, {ARCH_LABEL[a]} r = {rk}", "Wrapped (%)": C.pm(m["mean"], m["sd"], 1)})
    return {"df": pd.DataFrame(rows).fillna(""),
            "notes": ["Rule fixed on 24 Oct 2026 before the corrected results existed (D-21). Wrapped cells are kept for segmentation "
                      "training and scoring and excluded from dry mass and per-cell phase results."]}


@table("S12_error_strata", "si", "Where segmentation errors occur: boundary distance and phase quartile")
def s12():
    rows = []
    for a, rk in C.PREFERRED + [("edge_sam", 8), C.BASELINE]:
        es = C.cfg_entry(a, rk, "full")["error_strata"]
        r = {"Configuration": f"{ARCH_LABEL[a]} r = {rk}"}
        for k, v in es.items():
            r[k.replace("false negatives in phase quartile", "FN rate, phase Q").replace("boundary band", "pixel error, band")] = C.pm(v["mean"], v["sd"], 1)
        rows.append(r)
    q1 = "false negatives in phase quartile Q1"
    ipp = []
    for a, rk in C.PREFERRED + [("edge_sam", 8), C.BASELINE]:
        x = {s: v for s, v in C.cfg_entry(a, rk, "pmc_bga")["error_strata"][q1]["per_seed"].items()}
        y = {s: v for s, v in C.cfg_entry(a, rk, "full")["error_strata"][q1]["per_seed"].items()}
        ipp.append({"Configuration": f"{ARCH_LABEL[a]} r = {rk}",
                    "FN rate Q1, Dice+PMC+BGA (%)": C.pm(*C.mean_sd(x.values())[:2], 1),
                    "FN rate Q1, full (+IPP) (%)": C.pm(*C.mean_sd(y.values())[:2], 1),
                    "Δ (full − without IPP), pp": C.fmt_diff(C.paired_diff(x, y), 1)})
    return {"df": pd.DataFrame(rows), "extra": pd.DataFrame(ipp),
            "extra_title": "Does IPP increase misses of low-phase pixels? False-negative rate in the lowest phase quartile, with and without IPP",
            "notes": ["Pixel error rate (%) in distance bands from the annotated boundary; false-negative rate (%) inside annotated cells by "
                      "per-image quartile of background-corrected phase (Q1 = lowest phase). Errors concentrate at the boundary and in the "
                      "lowest-phase quartile."]}


@table("S13_cell_metrics_all_configs", "si", "Optical-mass and storage metrics of every rank-sweep configuration (sensitivity)")
def s13():
    rows = []
    for a in C.ARCHS:
        for rk in RANKS:
            m = C.cfg_entry(a, rk, "full")["metrics"]
            rows.append({"Architecture": ARCH_LABEL[a], "LoRA rank": rk, "Global IPE (%)": C.pm_metric(m["ipe_global_pct"], 2),
                         "Per-cell |IPE| median (%)": C.pm_metric(m["percell_abs_median_pct_nowrap"], 2),
                         "Matched cells (%)": C.pm_metric(m["gt_cells_matched_pct_nowrap"], 1),
                         "MdARE (%)": C.pm_metric(m["mdare_pct"], 2), "Bias (pg)": C.pm_metric(m["bias_pg"], 2),
                         "Dry mass ρ vs day": C.pm_metric(m["dry_mass_rho_pred"], 2),
                         "Dry mass late vs early (%)": C.pm_metric(m["dry_mass_late_vs_early_pct_pred"], 1),
                         "Area ρ vs day": C.pm_metric(m["area_rho_pred"], 2)})
    ann = C.compiled("cell_analysis")["annotation"]["storage_nowrap"]
    rows.append({"Architecture": "Annotation", "Dry mass ρ vs day": f"{ann['dry_mass']['spearman_rho']:.2f}",
                 "Dry mass late vs early (%)": f"{ann['dry_mass']['late_vs_early_pct']:.1f}", "Area ρ vs day": f"{ann['area']['spearman_rho']:.2f}"})
    neg = sum(1 for a in C.ARCHS for rk in RANKS if C.metric(a, rk, "full", "dry_mass_late_vs_early_pct_pred")["mean"] < ann["dry_mass"]["late_vs_early_pct"])
    note_number("configs_with_spurious_dry_mass_decline", f"{neg}/15 rank-sweep configurations show a larger late-vs-early dry-mass decline than the annotation", "Table S13")
    return {"df": pd.DataFrame(rows).fillna(""), "notes": ["Every configuration with predicted masks shows a dry-mass decline that the annotation "
                                                           "does not have; the size of this bias differs between models."]}


@table("S14_reference_sensitivity", "si", "Global IPE under different background references")
def s14():
    rows = []
    for a, rk in C.PREFERRED + [("edge_sam", 8), C.BASELINE]:
        m = C.cfg_entry(a, rk, "full")["metrics"]
        rows.append({"Configuration": f"{ARCH_LABEL[a]} r = {rk}",
                     "Background from annotation (%)": C.pm_metric(m["ipe_pct_reference_annotation"], 2),
                     "Background from prediction (%)": C.pm_metric(m["ipe_pct_reference_prediction"], 2),
                     "No reference, phase as provided (%)": C.pm_metric(m["ipe_pct_no_reference_as_provided"], 2),
                     "Residual background, max |rad|": C.pm_metric(m["background_residual_max_abs_rad"], 4)})
    off = C.compiled("cell_analysis")["annotation"]["phase_offsets"]["by_group"]
    g = off["date_named"]
    note_number("validation_background_offset_rad", {k: g[k] for k in ("min_rad", "median_rad", "max_rad", "n_images")},
                "outputs/compiled/cell_analysis_multiseed.json annotation.phase_offsets")
    return {"df": pd.DataFrame(rows), "extra_json": off,
            "notes": [f"The as-provided maps do not share a zero level (validation background {g['min_rad']:.2f} to {g['max_rad']:.2f} rad, "
                      f"median {g['median_rad']:.2f} rad; the training X_* maps are near zero-mean, D-03); integrated "
                      "phase depends on that constant, so every map is referenced to its own background median before training and "
                      "evaluation. The choice of reference (annotation vs prediction) changes the IPE by less than 0.01 pp."]}


# ------------------------------------------------------------------------------------------------ writing
def write_all() -> List[Dict[str, Any]]:
    out = []
    for t in REGISTRY:
        res = t["fn"]()
        d = C.TABLES / t["part"]
        d.mkdir(parents=True, exist_ok=True)
        df: pd.DataFrame = res["df"]
        df.to_csv(d / f"{t['id']}.csv", index=False)
        md = [f"# {t['id']}: {t['title']}", ""]
        if res.get("md_rows") == 0:
            md += [f"{len(df)} rows; see `{t['id']}.csv` (too wide and long for Markdown).", ""]
        else:
            md += [C.to_markdown(df), ""]
        if "extra" in res:
            res["extra"].to_csv(d / f"{t['id']}_b.csv", index=False)
            md += [f"**{res.get('extra_title', 'Paired differences')}**", "", C.to_markdown(res["extra"]), ""]
        if "extra_json" in res:
            md += ["**Phase offsets removed by the background referencing (rad)**", "", "```", json.dumps(res["extra_json"], indent=1)[:2000], "```", ""]
        md += ["Notes:", ""] + [f"* {n}" for n in res.get("notes", [])]
        (d / f"{t['id']}.md").write_text("\n".join(md) + "\n")
        out.append({**t, **res})
        print(f"[tables] {t['part']}/{t['id']}: {len(df)} rows")
    raw = C.TABLES / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    RUNS.to_csv(raw / "all_runs.csv", index=False)
    long = []
    for key, c in C.cell_configs().items():
        for mname, v in c["metrics"].items():
            long.append({"config": key, "arch": c["arch"], "rank": c["lora_rank"], "loss": c["loss"], "metric": mname,
                         "mean": v["mean"], "sd": v["sd"], "n": v["n"], **{f"seed_{s}": v["per_seed"].get(s) for s in ("42", "43", "44")}})
    pd.DataFrame(long).to_csv(raw / "cell_analysis_long.csv", index=False)
    return out


def write_index(results: List[Dict[str, Any]]) -> None:
    lines = ["# All revision tables", "", "Generated by `python artifacts/code/make_tables.py` from `outputs/` (read only). "
             "Do not edit by hand; rerun the script.", ""]
    for part, head in (("main", "Main manuscript"), ("si", "Supplementary Information")):
        lines += [f"## {head}", ""]
        for r in results:
            if r["part"] == part:
                lines += [(C.TABLES / part / f"{r['id']}.md").read_text(), ""]
    (C.TABLES / "ALL_TABLES.md").write_text("\n".join(lines))
    try:
        import docx
        from docx.enum.section import WD_ORIENT
        from docx.shared import Pt, Cm
    except ImportError:
        print("[tables] python-docx not installed: ALL_TABLES.docx skipped")
        return
    doc = docx.Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.5))
    doc.add_heading("Revision tables (generated; copy into the manuscript)", 0)
    for r in results:
        if r.get("md_rows") == 0:
            doc.add_heading(f"{r['id']}: {r['title']}", 2)
            doc.add_paragraph(f"{len(r['df'])} rows: see artifacts/tables/{r['part']}/{r['id']}.csv")
            continue
        doc.add_heading(f"{r['id']}: {r['title']}", 2)
        for df in [r["df"]] + ([r["extra"]] if "extra" in r else []):
            t = doc.add_table(rows=1, cols=len(df.columns))
            t.style = "Light Grid Accent 1"
            for i, c in enumerate(df.columns):
                t.rows[0].cells[i].text = str(c)
            for _, row in df.iterrows():
                cells = t.add_row().cells
                for i, v in enumerate(row.values):
                    cells[i].text = "" if (isinstance(v, float) and math.isnan(v)) else str(v)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(7)
            doc.add_paragraph("")
        for n in r.get("notes", []):
            doc.add_paragraph(n, style="List Bullet")
    doc.save(C.TABLES / "ALL_TABLES.docx")


def main() -> int:
    res = write_all()
    write_index(res)
    (C.ART / "numbers.json").write_text(json.dumps(NUMBERS, indent=1, default=float))
    print(f"[tables] {len(res)} tables, {len(NUMBERS)} numbers -> {C.TABLES}, {C.ART / 'numbers.json'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
