"""Small helpers to edit the manuscript .docx in place (python-docx + lxml), keeping the original paragraph formatting.

Inline markup in new text:  *italic*   ^superscript^   ~subscript~   **bold**.  "(r=8)" is written with an italic r.
Every run created here carries a highlight (yellow by default) so revised content is visible in Word.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path
from typing import List, Optional, Sequence

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu
from PIL import Image

HL = {"color": "yellow"}
TOKEN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\^[^^]+\^|~[^~]+~)")
DROP = {qn("w:r"), qn("w:proofErr"), qn("w:hyperlink"), qn("w:ins"), qn("w:del"), qn("w:smartTag"),
        "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath",
        "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMathPara"}


def italic_r(text: str) -> str:
    return re.sub(r"\(r=(\d+)\)", r"(*r*=\1)", text)


def base_rpr(p) -> OxmlElement:
    """Run properties to start from: first text run of the paragraph, else the paragraph mark's rPr."""
    for r in p.iter(qn("w:r")):
        if r.find(qn("w:t")) is not None and r.find(qn("w:rPr")) is not None:
            rpr = copy.deepcopy(r.find(qn("w:rPr")))
            break
    else:
        ppr = p.find(qn("w:pPr"))
        mark = ppr.find(qn("w:rPr")) if ppr is not None else None
        rpr = copy.deepcopy(mark) if mark is not None else OxmlElement("w:rPr")
        rpr.tag = qn("w:rPr")
    for tag in ("w:highlight", "w:i", "w:iCs", "w:b", "w:bCs", "w:vertAlign", "w:noProof", "w:lang"):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    return rpr


def _set(rpr, tag, val=None):
    el = OxmlElement(tag)
    if val is not None:
        el.set(qn("w:val"), val)
    # schema order: keep it simple and valid by inserting before sz/szCs/highlight where needed
    order = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike", "w:dstrike",
             "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden",
             "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
             "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout",
             "w:specVanish", "w:oMath"]
    pos = order.index(tag)
    for i, child in enumerate(list(rpr)):
        name = "w:" + child.tag.split("}")[1]
        if name in order and order.index(name) > pos:
            child.addprevious(el)
            return
    rpr.append(el)


def make_run(text: str, rpr, italic=False, bold=False, va=None, highlight=True):
    r = OxmlElement("w:r")
    rp = copy.deepcopy(rpr)
    if bold:
        _set(rp, "w:b")
        _set(rp, "w:bCs")
    if italic:
        _set(rp, "w:i")
        _set(rp, "w:iCs")
    if highlight and HL["color"]:
        _set(rp, "w:highlight", HL["color"])
    if va:
        _set(rp, "w:vertAlign", va)
    r.append(rp)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def runs_from_markup(text: str, rpr, highlight=True, bold_all=False) -> List:
    out = []
    for part in TOKEN.split(italic_r(text)):
        if not part:
            continue
        if part.startswith("**"):
            out.append(make_run(part[2:-2], rpr, bold=True, highlight=highlight))
        elif part.startswith("*"):
            out.append(make_run(part[1:-1], rpr, italic=True, bold=bold_all, highlight=highlight))
        elif part.startswith("^"):
            out.append(make_run(part[1:-1], rpr, va="superscript", bold=bold_all, highlight=highlight))
        elif part.startswith("~"):
            out.append(make_run(part[1:-1], rpr, va="subscript", bold=bold_all, highlight=highlight))
        else:
            out.append(make_run(part, rpr, bold=bold_all, highlight=highlight))
    return out


def set_text(p, text: str, highlight=True, rpr=None, prefix: Sequence = ()):
    """Replace the visible content of paragraph p (runs, equations) by text; pPr and bookmarks are kept.
    prefix: (text, bold) pieces written before text (e.g. a bold "TABLE 1" label)."""
    rpr = rpr if rpr is not None else base_rpr(p)
    for child in list(p):
        if child.tag in DROP:
            p.remove(child)
    anchor = None
    for child in p:
        if child.tag == qn("w:bookmarkEnd"):
            anchor = child
            break
    new = [make_run(t, rpr, bold=b, highlight=highlight) for t, b in prefix] + runs_from_markup(text, rpr, highlight)
    for r in new:
        if anchor is not None:
            anchor.addprevious(r)
        else:
            p.append(r)
    return p


def para_text(p) -> str:
    out = []
    for node in p.iter():
        if node.tag == qn("w:t") and node.text:
            out.append(node.text)
    return "".join(out)


def clone_par(template, text: str, highlight=True, prefix=(), keep_drawing=False):
    """New paragraph with the template's pPr and run formatting."""
    p = copy.deepcopy(template)
    for child in list(p):
        if child.tag not in (qn("w:pPr"),):
            p.remove(child)
    set_text(p, text, highlight, rpr=base_rpr(template), prefix=prefix)
    return p


def replace_image(doc, p, path: Path, width_in: Optional[float] = None, drop_svg=True):
    """Swap the picture inside paragraph p for the image at path (same rId target replaced), keep or set the width."""
    blip = next(p.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
    rid = blip.get(qn("r:embed"))
    part = doc.part.related_parts[rid]
    data = Path(path).read_bytes()
    part._blob = data
    if drop_svg:
        ext = blip.find("{http://schemas.openxmlformats.org/drawingml/2006/main}extLst")
        if ext is not None:
            blip.remove(ext)
    w, h = Image.open(path).size
    extent = next(p.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"))
    cx = int(width_in * 914400) if width_in else int(extent.get("cx"))
    cy = int(cx * h / w)
    extent.set("cx", str(cx))
    extent.set("cy", str(cy))
    for e in p.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}ext"):
        if e.get("cx"):
            e.set("cx", str(cx))
            e.set("cy", str(cy))
    return rid


def picture_par(doc, template_p, path: Path, width_in: float):
    """New picture paragraph (template paragraph formatting) holding a newly added image part."""
    p = copy.deepcopy(template_p)
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    tmp = doc.add_paragraph()
    run = tmp.add_run()
    run.add_picture(str(path), width=Emu(int(width_in * 914400)))
    r = tmp._p.find(qn("w:r"))
    p.append(r)
    tmp._p.getparent().remove(tmp._p)
    return p


# ------------------------------------------------------------------------------------------------ tables
def _cell(text, width_pct, bold, size, align, shade=None, highlight=True, borders=None):
    tc = OxmlElement("w:tc")
    tcpr = OxmlElement("w:tcPr")
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(int(width_pct * 50)))
    w.set(qn("w:type"), "pct")
    tcpr.append(w)
    if borders:
        tcpr.append(borders())
    if shade:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), shade)
        tcpr.append(shd)
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcpr.append(va)
    tc.append(tcpr)
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), "240")
    sp.set(qn("w:lineRule"), "auto")
    ppr.append(sp)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), align)
    ppr.append(jc)
    p.append(ppr)
    rpr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        f.set(qn(f"w:{a}"), "Times New Roman")
    rpr.append(f)
    for tag in ("w:sz", "w:szCs"):
        s = OxmlElement(tag)
        s.set(qn("w:val"), str(size))
        rpr.append(s)
    for line_no, line in enumerate(str(text).split("\n")):
        if line_no:
            br = OxmlElement("w:r")
            br.append(copy.deepcopy(rpr))
            br.append(OxmlElement("w:br"))
            p.append(br)
        for r in runs_from_markup(line, rpr, highlight=highlight, bold_all=bold):
            p.append(r)
    tc.append(p)
    return tc


def build_table(header: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float], style="main",
                highlight=True, shade_rows: Sequence[int] = (), group_rule_rows: Sequence[int] = ()):
    """style 'main': full grid, 8.5 pt (as Tables 1-7 of the final-review manuscript);
    style 'si': horizontal rules only, shaded header, 8 pt (as the SI tables). widths in percent of the text width."""
    size = 17 if style == "main" else 16
    tbl = OxmlElement("w:tbl")
    tblpr = OxmlElement("w:tblPr")
    st = OxmlElement("w:tblStyle")
    st.set(qn("w:val"), "PlainTable1" if style == "main" else "TableGrid")
    tblpr.append(st)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), "5000")
    tw.set(qn("w:type"), "pct")
    tblpr.append(tw)
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    tblpr.append(jc)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        none = style == "si" and side in ("left", "right", "insideV")
        e.set(qn("w:val"), "none" if none else "single")
        e.set(qn("w:sz"), "0" if none else "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "auto")
        b.append(e)
    tblpr.append(b)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), "57")
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)
    look = OxmlElement("w:tblLook")
    look.set(qn("w:val"), "04A0")
    look.set(qn("w:firstRow"), "1")
    look.set(qn("w:lastRow"), "0")
    look.set(qn("w:firstColumn"), "0")
    look.set(qn("w:lastColumn"), "0")
    look.set(qn("w:noHBand"), "1")
    look.set(qn("w:noVBand"), "1")
    tblpr.append(look)
    tbl.append(tblpr)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        g = OxmlElement("w:gridCol")
        g.set(qn("w:w"), str(int(w / 100 * 10454)))
        grid.append(g)
    tbl.append(grid)
    for ri, row in enumerate([header] + list(rows)):
        tr = OxmlElement("w:tr")
        trpr = OxmlElement("w:trPr")
        trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        tr.append(trpr)
        shade = "E7E6E6" if (ri == 0 and style == "si") else ("F2F2F2" if (ri - 1) in shade_rows and ri > 0 else None)
        for ci, (txt, w) in enumerate(zip(row, widths)):
            tr.append(_cell(txt, w, bold=(ri == 0), size=size, align="left" if ci == 0 else "center", shade=shade,
                            highlight=highlight))
        tbl.append(tr)
    return tbl
